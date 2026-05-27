import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def markdown_to_docx(md_path, docx_path):
    doc = Document()
    
    # Estilo base
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    if not os.path.exists(md_path):
        print(f"Error: No se encuentra {md_path}")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        if not line:
            continue

        if line.startswith('# '):
            h = doc.add_heading(line[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
        elif line.startswith('---'):
            doc.add_page_break()
        elif line.startswith('* ') or line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
            doc.add_paragraph(line[3:], style='List Number')
        elif line.startswith('```'):
            continue # Omitir tags de código por ahora o poner en courier
        else:
            p = doc.add_paragraph(line)

    doc.save(docx_path)
    print(f"Documento guardado en: {docx_path}")

if __name__ == "__main__":
    md_file = r'c:\Users\Sebastian\Documents\gymetra\GYMETRA-V1\doc\registration_documents\DNDA_Consolidado_Final.md'
    docx_file = r'c:\Users\Sebastian\Documents\gymetra\GYMETRA-V1\doc\registration_documents\REGISTRO_DNDA_GYMETRA_2026.docx'
    markdown_to_docx(md_file, docx_file)
